from __future__ import annotations

import json
import sys
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_paragraph(document: Document, text: str, indent: bool = True) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.45
    paragraph.paragraph_format.space_after = Pt(6)
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    run = paragraph.add_run(text)
    set_font(run, 11)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(level=level)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, 16 if level == 1 else 13, True)


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            set_font(run, 10, True)
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            if row_index % 2:
                shade(cells[index], "EAF2F8")
            for run in cells[index].paragraphs[0].runs:
                set_font(run, 9.5)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    document.add_paragraph()


def add_figure(document: Document, image: Path, caption: str, width_cm: float = 14.5) -> None:
    if not image.is_file():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_paragraph.add_run(caption)
    set_font(run, 9)


def main() -> None:
    results = ROOT / "results"
    metrics = json.loads((results / "metrics.json").read_text(encoding="utf-8"))
    template = ROOT / "resources" / "report_template.docx"
    if template.is_file():
        update_report(template, metrics)
        return
    raise FileNotFoundError(f"缺少正式报告模板: {template}")


def update_report(template: Path, metrics: dict) -> None:
    """Preserve the supplied report's cover, prose and figures; add regression."""
    document = Document(template)
    reg = metrics["regression"]
    official, grouped = reg["official_split"], reg["grouped_validation"]

    def replace(paragraph, content):
        props = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
        paragraph.clear()
        run = paragraph.add_run(content)
        if props is not None:
            run._r.insert(0, props)

    additions = {
        "本项目面向 CUB-200-2011": "在分类支路之外，系统以同一组图像特征独立回归体重、跗蹠长度、翅长、尾长及喙部三个尺寸，回归结果不依赖预测鸟种。",
        "官方测试集中共有": f"七项性状在官方图像留出测试中的平均 R² 为 {official['mean_r2']:.4f}，训练集内按鸟种分组的五折验证平均 R² 为 {grouped['mean_r2']:.4f}。",
        "程序采用 Python": "单图界面分别标注图像直接回归值和按预测鸟种查表的参考值。批处理增加七个带单位的回归列，统计页显示 MAE、RMSE 与 R²。",
        "AnalysisResult 是": "新增 regressed_traits 和 regression_note 字段；原 ecological_traits 保留查表结果。CSV 同时保存回归 JSON 和 regression_体重_g 等数值列。",
        "启动界面后": "回归结果位于单图结果区顶部，使用时可与后面的查表参考值对照。分类置信度仅针对分类输出，不是回归置信度。",
        "项目完成了从 CUB": "新增的七输出回归模型直接建立图像与连续形态性状之间的统计联系，并给出独立的测试预测明细、均值基线及回归消融结果。",
        "复杂背景、鸟体占比小": "回归精度有限，照片缺乏真实尺度与姿态控制，体型相似和拍摄背景都可能形成混杂。同一鸟种的各图共享性状标签，官方图像划分不能代替未见鸟种验证，也不能支持适应性演化或因果机制结论。",
    }
    for p in document.paragraphs:
        for prefix, extra in additions.items():
            if p.text.startswith(prefix):
                replace(p, p.text + extra)
                break
        if p.text.startswith("关键词："):
            replace(p, p.text + "；多输出回归")
        if p.text == "2.3 生态指标输出":
            replace(p, "2.3 按预测鸟种查询生态指标")

    def insert_before(title, build):
        anchor = next(p for p in document.paragraphs if p.text == title)._p
        previous = set(document._element.body)
        build()
        for element in list(document._element.body):
            if element not in previous and element.tag != qn("w:sectPr"):
                anchor.addprevious(element)

    def methods():
        add_heading(document, "2.4 图像与鸟类性状的直接回归", 2)
        add_paragraph(document, "设每张图像的 68 维视觉特征为 x，以图像真实类别关联的 AVONET 物种平均性状为监督标签 y。目标包含体重（g）、跗蹠长度、翅长、尾长、喙长、喙宽与喙深（mm）。这七列在 200 个类别中均完整，因此使用全部 11,788 张图像，没有删除样本或填造标签。核查时将 Vesper Sparrow 的错误学名修正为 Pooecetes gramineus，200 类现对应 200 个不同 AVONET 物种。栖息地和食性是类别变量，不作为连续回归目标；分布范围含缺失值，本次仍作为查表信息。")
        add_paragraph(document, "模型为 ExtraTreesRegressor 多输出回归器，共 500 棵树，max_features=1.0、max_leaf_nodes=256、min_samples_leaf=2，随机种子固定为 2026。为避免克与毫米的量纲差异主导树分裂，TransformedTargetRegressor 对训练标签逐列拟合 StandardScaler，预测时还原原单位。模型只读取颜色、形状、纹理、分形、对称性和分割质量；类别编号仅用于构造监督标签，文件名、分类预测、官方目标框与 IoU 均不作为回归输入。")
        add_paragraph(document, "部署模型仅由官方 5,994 张训练图像拟合，5,794 张测试图像仅用于最终评估。另在官方训练集内以 class_id 做 GroupKFold 五折验证，每折留出的鸟种与拟合鸟种互斥，覆盖全部 200 类。每折独立拟合目标标准化器和回归器，组外预测只用于诊断，不用于调参或替换部署模型。")
    insert_before("3 系统实现", methods)

    def findings():
        add_heading(document, "4.1 连续性状回归评估", 2)
        add_paragraph(document, "MAE 为平均绝对误差，RMSE 为均方根误差，均按目标原单位报告；R² 为决定系数。基线对所有测试图预测训练集的性状均值，不读取测试标签拟合。不同单位的误差不直接取总平均，整体结果仅报告七项 R² 的等权平均。下表的分组 R² 来自训练集内留鸟种的组外预测，与官方测试集评价范围不同。")
        add_table(document, ["性状", "MAE", "RMSE", "R²", "基线 RMSE", "分组 R²"],
            [[r["trait"].rsplit("_", 1)[0], f"{r['mae']:.3f}", f"{r['rmse']:.3f}", f"{r['r2']:.4f}", f"{r['baseline_rmse']:.3f}", f"{g['r2']:.4f}"]
             for r, g in zip(official["per_trait"], grouped["per_trait"])])
        add_paragraph(document, f"官方测试集平均 R² 为 {official['mean_r2']:.4f}。翅长的 R² 为 {official['per_trait'][2]['r2']:.4f}，体重为 {official['per_trait'][0]['r2']:.4f}，说明这些传统图像特征只解释了部分性状差异。全部七项 RMSE 均低于训练均值基线，但误差仍较大。留鸟种验证平均 R² 降至 {grouped['mean_r2']:.4f}，因此不能将同鸟种新照片上的效果直接外推到未见鸟种。")
        ablation = reg["ablation"]
        if ablation:
            groups = list(dict.fromkeys(r["feature_group"] for r in ablation))
            add_table(document, ["回归消融", "维数", "树数", "七性状平均 R²"],
                [[name, next(r['feature_count'] for r in ablation if r['feature_group'] == name), 150,
                  f"{sum(r['r2'] for r in ablation if r['feature_group'] == name) / 7:.4f}"] for name in groups])
            add_paragraph(document, "回归消融保持同一官方划分和训练标签标准化方案，三组模型均为 150 棵树。各性状 MAE、RMSE 和 R² 保存在 regression_ablation_results.csv。是否加入纹理、分形和对称性对误差的影响应以这些结果为依据，不能由分类准确率替代回归评价。")
        add_figure(document, ROOT / "results" / "regression_r2.png", "图 4 官方图像留出与训练集内留鸟种验证的回归 R²")
        add_figure(document, ROOT / "results" / "regression_predictions.png", "图 5 七项性状的官方测试集参考值与回归值")
        add_heading(document, "4.2 图像特征与性状的统计关联", 2)
        add_paragraph(document, "先按鸟种汇总官方训练图像的特征均值，再对 200 个鸟种的特征均值与 AVONET 性状计算 Spearman 相关。这样可避免把同一物种在多张照片上的重复标签当作独立生态观测。图中显示绝对相关较高的 18 项图像特征，完整矩阵保存于 regression_feature_trait_correlations.csv。颜色、图像中鸟体占比与性状的相关可能同时反映形态差异和摄影条件，只作统计解释，不构成适应性或因果证据。")
        add_figure(document, ROOT / "results" / "regression_correlations.png", "图 6 训练鸟种均值的图像特征与形态性状相关")
    insert_before("5 界面与使用流程", findings)
    # Replace the existing screenshot without adding a second copy.
    for i, p in enumerate(document.paragraphs):
        old_figures = {"图 1  200 个鸟种的全特征 PCA 分布": "pca_species_features.png",
                       "图 2  ExtraTrees 的前 20 个特征重要性": "feature_importance.png",
                       "图 3  完整测试集混淆矩阵": "confusion_matrix.png"}
        if p.text in old_figures:
            image_p = document.paragraphs[i - 1]
            image_p.clear()
            image_p.add_run().add_picture(str(ROOT / "results" / old_figures[p.text]), width=Cm(14.5))
        if p.text == "图 4  系统桌面界面":
            image_p = document.paragraphs[i - 1]
            image_p.clear()
            image_p.add_run().add_picture(str(ROOT / "results" / "gui_screenshot.png"), width=Cm(14.5))
            replace(p, "图 7 系统桌面界面及直接回归输出")
    output = ROOT / "docs" / "鸟类生态图像评价系统课程报告.docx"
    add_paragraph(document, "Cornell Lab of Ornithology. Vesper Sparrow (Pooecetes gramineus). All About Birds. https://www.allaboutbirds.org/guide/Vesper_Sparrow/overview （学名核对，2026-09-16）", indent=False)
    for p in document.paragraphs:
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
        if p.text.startswith("图 "):
            p.paragraph_format.keep_together = True
        if p.text.startswith(("2.4 ", "4.1 ")):
            p.paragraph_format.page_break_before = True
    for table in document.tables[1:]:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
    document.save(output)
    print(output)


if __name__ == "__main__":
    main()
